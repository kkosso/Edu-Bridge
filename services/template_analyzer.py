"""
services/template_analyzer.py
==============================
양식을 셀 단위로 분석하고 Gemini로 각 셀의 의미를 파악
"""
import json
from typing import List, Dict
from docx import Document


def extract_template_cells(file_path: str) -> List[Dict]:
    """
    양식의 모든 셀을 위치 정보와 함께 추출 (병합 셀 처리)

    Returns:
        [
            {"table_idx": 0, "row": 0, "col": 0, "text": "단원(차시)", "is_empty": False},
            {"table_idx": 0, "row": 0, "col": 1, "text": "", "is_empty": True},
            ...
        ]
    """
    doc = Document(file_path)
    cells_info = []

    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            seen_in_row = set()  # 행 내 병합 셀 추적
            for ci, cell in enumerate(row.cells):
                text = cell.text.strip()
                # 병합 셀: 같은 셀 객체가 반복되면 첫 번째만 사용
                cell_key = id(cell._tc)
                if cell_key in seen_in_row:
                    continue
                seen_in_row.add(cell_key)

                cells_info.append({
                    "table_idx": ti,
                    "row": ri,
                    "col": ci,
                    "text": text[:200],  # 너무 긴 텍스트는 자르기
                    "is_empty": not text or _is_placeholder(text),
                })

    return cells_info


def _is_placeholder(text: str) -> bool:
    """예시/플레이스홀더 텍스트인지 판단"""
    if not text:
        return True
    text = text.strip()
    if len(text) < 2:
        return True
    placeholders = ["...", "___", "OOO", "○○○", "(예", "예)", "예시"]
    return any(p in text for p in placeholders)


def analyze_template_deeply(file_path: str) -> Dict:
    """
    Gemini를 사용해 양식의 각 셀을 분석.

    Returns:
        {
            "cells": [
                {
                    "table_idx": 0, "row": 0, "col": 0,
                    "role": "label",  # "label" or "data"
                    "category": "lesson_meta",  # personal_info, lesson_meta, lesson_content, evaluation
                    "label_text": "단원(차시)",
                },
                {
                    "table_idx": 0, "row": 0, "col": 1,
                    "role": "data",
                    "for_label": "단원(차시)",
                    "category": "lesson_meta",
                    "expected_content": "단원의 이름과 차시 (예: 2단원 모양 알기 (3/8))",
                    "approx_length": "short",  # short / medium / long
                },
                ...
            ],
            "summary": "유아 활동 지도안 양식 - 표 6개로 구성. 개인정보, 수업 메타, 수업 내용, 평가 영역 포함",
        }
    """
    cells_info = extract_template_cells(file_path)
    if not cells_info:
        return {"cells": [], "summary": "표가 없는 양식 (단락 기반)"}

    # 셀 정보를 Gemini에게 보낼 텍스트로 정리
    cells_text = []
    for c in cells_info:
        marker = "[비어있음]" if c["is_empty"] else c["text"]
        cells_text.append(f"표{c['table_idx']} 행{c['row']} 열{c['col']}: {marker}")

    cells_str = "\n".join(cells_text)

    prompt = f"""아래는 한국 유아교육 지도안 양식의 모든 셀 정보입니다.
각 셀을 분석해서 역할과 의미를 파악해주세요.

## 셀 목록
{cells_str}

## 작업 지침
각 셀에 대해 다음을 판단하세요:

1. **role**: "label" (헤더/라벨) 또는 "data" (값이 들어갈 셀)
2. **category** (label인 경우):
   - "personal_info": 교사명, 학생명, 학교명, 학과, 결재란 등 사람/조직 정보
   - "lesson_meta": 단원, 차시, 학습 주제, 학습 목표, 성취 기준, 교과 역량 등 수업 정보
   - "lesson_content": 학습 단계, 도입, 전개, 정리, 교사 활동, 학생 활동, 유의점 등 수업 내용
   - "evaluation": 평가 기준, 평가 방법, 성취 수준 등 평가 영역
   - "resource": 교수학습자료, 준비물 등
3. **for_label** (data인 경우): 어떤 라벨에 대한 데이터인지
4. **expected_content** (data인 경우): 어떤 내용이 들어가야 하는지 구체적으로 설명
5. **approx_length** (data인 경우): "short" (한 줄), "medium" (몇 문장), "long" (한 단락 이상)

## 출력 형식
오직 JSON만 출력. 다른 설명 없이:
{{
  "cells": [
    {{"table_idx": 0, "row": 0, "col": 0, "role": "label", "category": "lesson_meta", "label_text": "단원(차시)"}},
    {{"table_idx": 0, "row": 0, "col": 1, "role": "data", "for_label": "단원(차시)", "category": "lesson_meta", "expected_content": "단원 이름과 차시 표기 (예: 2단원 모양 알기 3/8차시)", "approx_length": "short"}},
    ...
  ],
  "summary": "이 양식에 대한 1-2줄 설명"
}}

모든 셀에 대해 분석하세요. 단, 같은 텍스트가 반복되는 병합 셀은 한 번만 분석.
빈 셀이라도 어떤 라벨의 데이터인지 추론해서 role을 "data"로 분류하고 for_label과 expected_content를 채워주세요.
"""

    try:
        from services.keyword_extractor import _get_client
        from google.genai import types as genai_types

        client = _get_client()
        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=prompt,
            config=genai_types.GenerateContentConfig(
                temperature=0.2,
                max_output_tokens=8192,
                response_mime_type="application/json",
                thinking_config=genai_types.ThinkingConfig(thinking_budget=0),
            ),
        )

        text = response.text.strip()
        if text.startswith("```"):
            text = text.split("```")[1]
            if text.startswith("json"):
                text = text[4:].strip()
            text = text.split("```")[0].strip()

        result = json.loads(text)
        if "cells" not in result:
            result["cells"] = []
        return result
    except Exception as e:
        print(f"⚠️ 양식 심층 분석 실패: {e}")
        # 실패해도 빈 결과 반환 (기본 채우기 폴백)
        return {"cells": [], "summary": f"분석 실패: {e}", "error": str(e)}


def fill_cells_with_lesson(
    template_analysis: Dict,
    lesson_markdown: str,
    lesson_meta: Dict = None,
    personal_info: Dict = None,
) -> Dict:
    """
    심층 분석된 양식 정보 + 지도안 마크다운 → 각 데이터 셀에 들어갈 내용 매핑

    Returns:
        {
            "(table_idx, row, col)": "채울 내용",
            ...
        }
    """
    cells = template_analysis.get("cells", [])
    data_cells = [c for c in cells if c.get("role") == "data"]

    if not data_cells:
        return {}

    # 개인정보 셀과 수업 내용 셀 분리
    # personal_info에 매칭되는 셀은 사용자 입력값 사용
    direct_fills = {}  # (ti,ri,ci) -> content
    cells_for_gemini = []

    for c in data_cells:
        position = (c["table_idx"], c["row"], c["col"])
        category = c.get("category", "")
        for_label = c.get("for_label", "")

        # 개인정보 셀: 사용자 입력에서 직접 매칭
        if category == "personal_info" and personal_info:
            matched = None
            for label, value in personal_info.items():
                if label in for_label or for_label in label:
                    matched = value
                    break
            if matched:
                direct_fills[position] = matched
                continue

        # 나머지는 Gemini에게 매핑 요청
        cells_for_gemini.append(c)

    # Gemini로 수업 내용 셀 채우기
    if cells_for_gemini:
        gemini_fills = _gemini_fill_data_cells(
            cells_for_gemini, lesson_markdown, lesson_meta or {}
        )
        for position, content in gemini_fills.items():
            if position not in direct_fills:
                direct_fills[position] = content

    return direct_fills


def _gemini_fill_data_cells(
    data_cells: List[Dict],
    lesson_markdown: str,
    lesson_meta: Dict,
) -> Dict:
    """데이터 셀들에 대해 Gemini로 내용 생성"""
    if not data_cells:
        return {}

    # 셀 목록을 ID와 함께 정리
    cells_for_prompt = []
    for i, c in enumerate(data_cells):
        cells_for_prompt.append({
            "id": i,
            "label": c.get("for_label", "(unknown)"),
            "category": c.get("category", ""),
            "expected": c.get("expected_content", ""),
            "length": c.get("approx_length", "medium"),
        })

    meta_str = ""
    if lesson_meta:
        parts = []
        if "age" in lesson_meta:
            parts.append(f"대상 연령: 만 {lesson_meta['age']}세")
        if "duration" in lesson_meta:
            parts.append(f"수업 시간: {lesson_meta['duration']}분")
        if "search_query" in lesson_meta:
            parts.append(f"수업 주제: {lesson_meta['search_query']}")
        if parts:
            meta_str = "\n## 수업 메타 정보\n" + "\n".join(parts)

    prompt = f"""아래는 [생성된 지도안]과 [채워야 할 양식 셀 목록]입니다.
각 셀에 들어갈 내용을 지도안 내용을 바탕으로 작성해주세요.

## 생성된 지도안 (마크다운)
{lesson_markdown}
{meta_str}

## 채워야 할 양식 셀 목록
{json.dumps(cells_for_prompt, ensure_ascii=False, indent=2)}

## 작업 지침
1. 각 셀의 label과 expected_content를 보고 지도안에서 적절한 내용을 추출/요약하세요.
2. length가 "short"면 한 줄 이내, "medium"이면 1-3 문장, "long"이면 한 단락으로.
3. 마크다운 기호(#, *, -)는 제거하고 깔끔한 문장으로.
4. 지도안에 직접적인 정보가 없는 경우, 수업 메타 정보와 expected_content 설명에 기반해 합리적으로 작성.
5. 평가/성취기준 등은 일반적인 유아교육 기준에 맞춰 적절히 생성.
6. 학습 단계(도입/전개/정리)별 셀이면 지도안의 해당 단계 내용을 정확히 매핑.
7. 모든 셀에 내용을 채워주세요. 빈 값 없이.

## 출력 형식
오직 JSON만 출력:
{{
  "fills": [
    {{"id": 0, "content": "..."}},
    {{"id": 1, "content": "..."}},
    ...
  ]
}}
"""

    try:
        from services.keyword_extractor import _get_client
        from google.genai import types as genai_types

        client = _get_client()
        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=prompt,
            config=genai_types.GenerateContentConfig(
                temperature=0.4,
                max_output_tokens=8192,
                response_mime_type="application/json",
                thinking_config=genai_types.ThinkingConfig(thinking_budget=0),
            ),
        )

        text = response.text.strip()
        if text.startswith("```"):
            text = text.split("```")[1]
            if text.startswith("json"):
                text = text[4:].strip()
            text = text.split("```")[0].strip()

        # JSON 파싱 시도 — 실패 시 자동 복구
        def try_parse(raw: str) -> dict:
            try:
                return json.loads(raw)
            except json.JSONDecodeError:
                pass

            # 1차 복구: 잘린 JSON 끝 처리 (마지막 완전한 fill 항목까지만 살리기)
            try:
                last_brace = raw.rfind('}')
                if last_brace != -1:
                    candidate = raw[:last_brace + 1]
                    # fills 배열이 닫히지 않은 경우 닫아주기
                    if '"fills"' in candidate and candidate.count('[') > candidate.count(']'):
                        candidate = candidate + ']}'
                    elif '"fills"' in candidate and not candidate.rstrip().endswith('}}'):
                        candidate = candidate + '}'
                    return json.loads(candidate)
            except Exception:
                pass

            # 2차 복구: fills 배열만 추출
            try:
                import re as _re
                m = _re.search(r'"fills"\s*:\s*(\[.*?\])', raw, _re.S)
                if m:
                    arr = json.loads(m.group(1))
                    return {"fills": arr}
            except Exception:
                pass

            # 3차 복구: 개별 {"id":N, "content":"..."} 패턴 추출
            try:
                import re as _re
                items = _re.findall(r'\{[^{}]*?"id"\s*:\s*(\d+)[^{}]*?"content"\s*:\s*"((?:[^"\\]|\\.)*)"\s*\}', raw)
                if items:
                    return {"fills": [{"id": int(i), "content": c} for i, c in items]}
            except Exception:
                pass

            raise ValueError(f"JSON 복구 실패. 원본 앞 200자: {raw[:200]}")

        result = try_parse(text)
        fills = result.get("fills", [])

        # id → (ti, ri, ci) 매핑
        position_fills = {}
        for fill in fills:
            idx = fill.get("id")
            content = fill.get("content", "")
            if idx is not None and 0 <= idx < len(data_cells) and content:
                c = data_cells[idx]
                position = (c["table_idx"], c["row"], c["col"])
                position_fills[position] = content

        return position_fills
    except Exception as e:
        print(f"⚠️ Gemini 셀 채우기 실패: {e}")
        return {}
